"""Section parser — extract text and sections from PDF/DOCX resumes."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Optional

from app.utils.logger import get_logger

logger = get_logger("section_parser")

# Section header patterns (case-insensitive)
SECTION_PATTERNS = {
    "contact": re.compile(r"(?i)^(?:contact|personal)\s*(?:info|information|details)?$"),
    "summary": re.compile(r"(?i)^(?:summary|objective|profile|about\s*me|overview)$"),
    "experience": re.compile(r"(?i)^(?:experience|work\s*(?:history|experience)|employment|professional\s*experience)$"),
    "education": re.compile(r"(?i)^(?:education|academic|qualifications|degrees?)$"),
    "skills": re.compile(r"(?i)^(?:skills|technical\s*skills|core\s*(?:competencies|skills)|technologies)$"),
    "projects": re.compile(r"(?i)^(?:projects|portfolio|personal\s*projects|key\s*projects)$"),
    "certifications": re.compile(r"(?i)^(?:certifications?|licenses?|accreditations?|credentials)$"),
}

MAGIC_PDF = b"%PDF"
MAGIC_DOCX = b"PK\x03\x04"


def validate_magic_bytes(content: bytes, extension: str) -> bool:
    """Validate file content against expected magic bytes."""
    if extension.lower() == ".pdf":
        return content[:4] == MAGIC_PDF
    if extension.lower() == ".docx":
        return content[:4] == MAGIC_DOCX
    return False


def extract_text_from_pdf(file_path: str) -> str:
    """Extract text from PDF using pdfplumber, fallback to OCR."""
    try:
        import pdfplumber
        text = ""
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
        if len(text.strip()) < 100:
            logger.info("parser.ocr_fallback", reason="insufficient_text", chars=len(text.strip()))
            text = _ocr_extract(file_path)
        return text
    except Exception as exc:
        logger.error("parser.pdf_failed", error=str(exc))
        return _ocr_extract(file_path)


def _ocr_extract(file_path: str) -> str:
    """OCR fallback using pytesseract."""
    try:
        from PIL import Image
        import pytesseract
        from pdf2image import convert_from_path
        images = convert_from_path(file_path)
        text = ""
        for img in images:
            text += pytesseract.image_to_string(img) + "\n"
        return text
    except ImportError:
        logger.warning("parser.ocr_unavailable", reason="pytesseract_not_installed")
        return ""
    except Exception as exc:
        logger.error("parser.ocr_failed", error=str(exc))
        return ""


def extract_text_from_docx(file_path: str) -> str:
    """Extract text from DOCX using python-docx."""
    try:
        import docx
        doc = docx.Document(file_path)
        parts = []
        for para in doc.paragraphs:
            parts.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return "\n".join(parts)
    except Exception as exc:
        logger.error("parser.docx_failed", error=str(exc))
        return ""


def detect_sections(text: str) -> dict[str, str]:
    """Parse text into named sections using header detection."""
    sections: dict[str, str] = {}
    current_section = "other"
    section_lines: dict[str, list[str]] = {"other": []}
    lines = text.split("\n")

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        matched = False
        for section_key, pattern in SECTION_PATTERNS.items():
            if pattern.match(stripped):
                current_section = section_key
                if section_key not in section_lines:
                    section_lines[section_key] = []
                matched = True
                break
        if not matched:
            if current_section not in section_lines:
                section_lines[current_section] = []
            section_lines[current_section].append(stripped)

    for key, lines_list in section_lines.items():
        content = "\n".join(lines_list).strip()
        if content:
            sections[key] = content

    return sections


def clean_text(text: str) -> str:
    """Clean and normalize extracted text."""
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"[•●◆▪–—]", " ", text)
    return text.strip()


def parse_resume(file_path: str, extension: Optional[str] = None) -> dict:
    """Main parse function — returns structured resume data.

    Returns: {
        "full_text": str,
        "sections": { "experience": str, "education": str, ... }
    }
    """
    path = Path(file_path)
    ext = extension or path.suffix.lower()

    if ext == ".pdf":
        raw_text = extract_text_from_pdf(file_path)
    elif ext == ".docx":
        raw_text = extract_text_from_docx(file_path)
    else:
        logger.warning("parser.unsupported_type", ext=ext)
        return {"full_text": "", "sections": {}}

    full_text = clean_text(raw_text)
    sections = detect_sections(raw_text)

    logger.info("parser.complete", chars=len(full_text), sections=list(sections.keys()))
    return {"full_text": full_text, "sections": sections}
